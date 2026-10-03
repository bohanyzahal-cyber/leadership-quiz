"""Native, editable navigation for the ten-page Hebrew exam handbook.

Navigation is kept separate from the verified teaching text. The Word importer
skips these styles/table captions and reads navigation.json for the HTML edition.
"""
from copy import deepcopy
from pathlib import Path
from functools import lru_cache
import json
import re

from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_DIRECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree

HERE = Path(__file__).resolve().parent
NAV = json.loads((HERE / 'navigation.json').read_text(encoding='utf-8'))
PREFIX = 'ExamNav'


def plain(text):
    return re.sub(r'[\u200e\u200f*]', '', text)


@lru_cache(maxsize=1)
def helpers():
    scope = {'__file__': str(HERE / 'to_docx.py'), '__name__': 'exam_formatting'}
    exec((HERE / 'to_docx.py').read_text(encoding='utf-8').split('\ndoc = Document()')[0], scope)
    return scope['rich'], scope['rtl']


def caption_table(table, value):
    caption = OxmlElement('w:tblCaption')
    caption.set(qn('w:val'), value)
    table._tbl.tblPr.append(caption)


def bookmark(paragraph, name, bookmark_id):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def page_link(paragraph, anchor, display_page, size=9.5):
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'), anchor)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGEREF ' + anchor + ' \\h')
    holder = Paragraph(OxmlElement('w:p'), paragraph._parent)
    rich, _ = helpers()
    rich(holder, str(display_page), size, True)
    for run in list(holder._p):
        if run.tag == qn('w:r'):
            field.append(run)
    link.append(field)
    paragraph._p.append(link)


def set_columns(section, count, gap=170):
    cols = section._sectPr.find(qn('w:cols'))
    for child in list(cols):
        cols.remove(child)
    cols.set(qn('w:num'), str(count))
    cols.set(qn('w:space'), str(gap))
    cols.set(qn('w:equalWidth'), '1')


def compact_section_separator(section, doc):
    parent=section._sectPr.getparent()
    if parent.tag == qn('w:pPr'):
        paragraph=Paragraph(parent.getparent(),doc)
        paragraph.paragraph_format.line_spacing=Pt(1)
        paragraph.paragraph_format.space_before=Pt(0)
        paragraph.paragraph_format.space_after=Pt(0)


def edge_tab(header, chapter_index, label):
    """Page-relative editable VML text box; no extra body line or image."""
    vml = 'urn:schemas-microsoft-com:vml'
    office = 'urn:schemas-microsoft-com:office:office'
    pict = OxmlElement('w:pict')
    shape = etree.SubElement(pict, '{' + vml + '}rect', nsmap={'v': vml, 'o': office})
    shape.set('id', PREFIX + 'Tab' + str(chapter_index))
    shape.set('style', 'position:absolute;left:5mm;top:' + str(20 + (chapter_index-1)*24.5) +
              'mm;width:5mm;height:24mm;z-index:251659264;'
              'mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
    shape.set('fillcolor', '#E8EDF1')
    shape.set('strokecolor', '#87919B')
    shape.set('strokeweight', '0.5pt')
    textbox = etree.SubElement(shape, '{' + vml + '}textbox')
    textbox.set('style', 'mso-layout-flow-alt:bottom-to-top')
    textbox.set('inset', '0.5mm,0.2mm,0.5mm,0.2mm')
    content = OxmlElement('w:txbxContent')
    textbox.append(content)
    p = Paragraph(OxmlElement('w:p'), header._parent)
    rich, rtl = helpers()
    rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    rich(p, str(chapter_index) + ' ' + label, 8.5, True)
    content.append(p._p)
    header.add_run()._r.append(pict)


def apply_navigation(doc, data):
    """Apply once to the existing Word source or a newly exported handbook."""
    if any(p.style.name == PREFIX + 'Cue' for p in doc.paragraphs):
        raise ValueError('Navigation already exists; edit its named paragraphs in place.')
    if len(data['sections']) != len(NAV['chapters']):
        raise ValueError('Navigation and handbook chapters must agree.')
    rich, rtl = helpers()
    for suffix, size, bold in [('Title',12,True),('Hint',9.5,False),('Cue',10,False),('Index',9.5,False)]:
        style = doc.styles.add_style(PREFIX + suffix, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:cs'), 'Arial')
        if bold:
            style._element.rPr.append(OxmlElement('w:bCs'))
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = Pt(11 if suffix == 'Index' else size*1.18)
        style.paragraph_format.keep_with_next = suffix == 'Title'

    headings = [next(p for p in doc.paragraphs if plain(p.text) == s['name']) for s in data['sections']]
    for i, (heading, chapter) in enumerate(zip(headings, NAV['chapters']),1):
        bookmark(heading, PREFIX+'Chapter'+str(i), 1000+i)
        cue = heading.insert_paragraph_before(style=PREFIX+'Cue')
        heading._p.addnext(cue._p)
        rtl(cue)
        cue.paragraph_format.space_after = Pt(5)
        cue.paragraph_format.keep_with_next = True
        key, _, explanation = chapter['cue'].partition(' :: ')
        rich(cue, key+' — '+explanation,10)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'),'F1F3F5')
        cue._p.get_or_add_pPr().append(shading)

    # Keep all teaching text and body font sizes. Sources are secondary text.
    for paragraph in doc.paragraphs:
        if paragraph.style.name == 'Heading 2':
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'),'E8EDF1')
            paragraph._p.get_or_add_pPr().append(shading)
        if plain(paragraph.text).startswith(('מקור:', 'מקורות:')):
            for run in paragraph.runs:
                run.font.size = Pt(9)
                for old in run._r.get_or_add_rPr().findall(qn('w:szCs')):
                    run._r.rPr.remove(old)
                size = OxmlElement('w:szCs'); size.set(qn('w:val'),'18'); run._r.rPr.append(size)
        if 'בר אילן   עדכון ' in paragraph.text:
            paragraph.clear(); rtl(paragraph)
            rich(paragraph,'בר אילן   עדכון '+NAV['updated']+'   5 דפים דו־צדדיים',8.5)

    # Comparison tables must be readable in one place, not split between the
    # bottom of one column and the top of the other with repeated headers.
    for table in doc.tables:
        for row_index,row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows)-1

    # Empty separator paragraphs after the last chapter's tables need only a
    # small gap. Their former body line height wastes space needed by the index.
    in_last=False
    for node in doc.element.body:
        if node is headings[-1]._p: in_last=True
        if in_last and node.tag==qn('w:p') and not node.xpath('.//w:t | .//w:br | .//w:sectPr'):
            previous=node.getprevious()
            if previous is not None and previous.tag==qn('w:tbl'):
                p=Paragraph(node,doc)
                p.paragraph_format.line_spacing=Pt(1)
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.space_after=Pt(3)

    # Compact table at the very beginning, rather than spending a page on a cover.
    first = headings[0]
    title = first.insert_paragraph_before(style=PREFIX+'Title')
    rtl(title); rich(title,'מפת עמודים',12,True)
    bookmark(title,PREFIX+'Map',1100)
    table = doc.add_table(rows=0,cols=2)
    caption_table(table,PREFIX+'Map')
    table.autofit = False
    table.table_direction = WD_TABLE_DIRECTION.RTL
    widths=[81,10]
    for col,width in zip(table.columns,widths): col.width=Mm(width)
    for i,chapter in enumerate(NAV['chapters'],1):
        row=table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.width=Mm(widths[j])
            pr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for key,value in [('top',12),('bottom',12),('left',30),('right',30)]:
                e=OxmlElement('w:'+key);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5' if i%2 else 'FFFFFF');pr.append(sh)
            p=cell.paragraphs[0];rtl(p)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(12)
            if j==0:
                rich(p,chapter['label'],10)
            else:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                page_link(p,PREFIX+'Chapter'+str(i),i,10)
    first._p.addprevious(table._tbl)
    hint=first.insert_paragraph_before(style=PREFIX+'Hint');rtl(hint)
    rich(hint,'לא זוכרים את הפרק? מפתח המונחים בעמוד ',9.5)
    page_link(hint,PREFIX+'Index',10)
    hint.paragraph_format.space_after=Pt(6)

    # The final chapter is already followed by a continuous balancing section.
    # Its spare bottom area holds a full-width heading and five index columns.
    set_columns(doc.sections[-1],1)
    title=doc.add_paragraph(style=PREFIX+'Title');rtl(title)
    rich(title,'מפתח מונחים',12,True)
    bookmark(title,PREFIX+'Index',1101)
    title.paragraph_format.space_before=Pt(3)
    hint=doc.add_paragraph(style=PREFIX+'Hint');rtl(hint)
    rich(hint,'עברית לפי א״ב ואחריה אנגלית. המספר מפנה לעמוד בחוברת.',9.5)
    hint.paragraph_format.space_after=Pt(3)
    index_section=doc.add_section(WD_SECTION.CONTINUOUS)
    compact_section_separator(doc.sections[-2],doc)
    set_columns(index_section,5)
    terms=[(term,i) for i,c in enumerate(NAV['chapters'],1) for term in c['terms']]
    terms.sort(key=lambda x:(bool(re.match(r'[A-Za-z]',x[0])),x[0]))
    for term,page in terms:
        p=doc.add_paragraph(style=PREFIX+'Index');rtl(p)
        rich(p,term+'\u00a0',9.5)
        page_link(p,PREFIX+'Chapter'+str(page),page)
    # End in the index's five columns. An empty one-column section after them
    # forces a blank eleventh page when the index reaches the bottom margin.

    for i,chapter in enumerate(NAV['chapters'],1):
        section=doc.sections[i-1]
        header=section.header.paragraphs[0]
        edge_tab(header,i,chapter['tab'])
        footer=section.footer.paragraphs[0]
        footer.clear();rtl(footer)
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rich(footer,'עמוד ',8.5)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
        rich(footer,' מתוך ',8.5)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'NUMPAGES');footer._p.append(field)
        rich(footer,'   ·   מפת עמודים ',8.5);page_link(footer,PREFIX+'Map',1,8.5)
        rich(footer,'   ·   מפתח מונחים ',8.5);page_link(footer,PREFIX+'Index',10,8.5)
    return {'chapters':len(headings),'index_terms':len(terms),'cues':len(headings)}
